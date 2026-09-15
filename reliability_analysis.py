"""
Assignment 1 - Reliability Engineering and Failure Analysis
Astana IT University | Fault Tolerance | 2026-2027
Student: Zhumabayev Magzhan (ID 255408)

This script reproduces every number used in the report:
  Part A - downtime / uptime / availability / MTTF / MTTR / MTBF
  Part B - Reliability Block Diagram (series-parallel) calculation
  Part C - FMEA with RPN = S x O x D
  Part D - Fault Tree (AND/OR) top-event probability
It also exports reliability_calculations.xlsx and the RBD/FTA diagrams (PNG).

Run:  python reliability_analysis.py
"""
from collections import defaultdict
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import Rectangle, FancyBboxPatch, Circle, Polygon
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# --------------------------------------------------------------------------
# Input data (synthetic, from the assignment)
# --------------------------------------------------------------------------
OPERATING_HOURS = 720  # 30 days

# (event, component, failure_time, repair_completed, duration)
EVENTS = [
    (1,  "Application Server A", 38,  40,  2),
    (2,  "Primary Database",     91,  95,  4),
    (3,  "Campus Network",       143, 144, 1),
    (4,  "Application Server B", 201, 203, 2),
    (5,  "Primary Database",     287, 291, 4),
    (6,  "Load Balancer",        356, 357, 1),
    (7,  "Application Server A", 411, 413, 2),
    (8,  "Campus Network",       478, 480, 2),
    (9,  "Primary Database",     529, 534, 5),
    (10, "Application Server B", 601, 603, 2),
    (11, "Application Server A", 654, 655, 1),
    (12, "Load Balancer",        689, 690, 1),
]

RELIABILITY = {
    "Load Balancer":        0.995,
    "Application Server A": 0.970,
    "Application Server B": 0.970,
    "Primary Database":     0.980,
    "Standby Database":     0.990,
    "Campus Network":       0.995,
}

SINGLE_POINTS_OF_FAILURE = {"Load Balancer", "Campus Network"}


def section(title):
    print("\n" + "=" * 72)
    print(title)
    print("=" * 72)


# --------------------------------------------------------------------------
# Part A - Reliability metrics
# --------------------------------------------------------------------------
section("PART A - RELIABILITY METRICS")

# sanity check: duration must equal repair_completed - failure_time
for ev in EVENTS:
    assert ev[3] - ev[2] == ev[4], f"Duration mismatch in event {ev[0]}"

n_failures = len(EVENTS)
total_downtime = sum(ev[4] for ev in EVENTS)
total_uptime = OPERATING_HOURS - total_downtime
availability = total_uptime / OPERATING_HOURS
mttr = total_downtime / n_failures
mttf = total_uptime / n_failures
mtbf = OPERATING_HOURS / n_failures          # = MTTF + MTTR

print(f"Number of failure events     : {n_failures}")
print(f"Total downtime               : {total_downtime} h  "
      f"({' + '.join(str(ev[4]) for ev in EVENTS)})")
print(f"Total uptime                 : {OPERATING_HOURS} - {total_downtime} = {total_uptime} h")
print(f"Availability                 : {total_uptime}/{OPERATING_HOURS} = {availability:.4f} "
      f"= {availability*100:.2f} %")
print(f"MTTF = uptime / failures     : {total_uptime}/{n_failures} = {mttf:.2f} h")
print(f"MTTR = downtime / failures   : {total_downtime}/{n_failures} = {mttr:.2f} h")
print(f"MTBF = period / failures     : {OPERATING_HOURS}/{n_failures} = {mtbf:.2f} h "
      f"(check: MTTF + MTTR = {mttf + mttr:.2f} h)")
print(f"Availability check A = MTTF/(MTTF+MTTR) = {mttf/(mttf+mttr):.4f}")

# per-component breakdown
per_comp = defaultdict(lambda: {"failures": 0, "downtime": 0})
for ev in EVENTS:
    per_comp[ev[1]]["failures"] += 1
    per_comp[ev[1]]["downtime"] += ev[4]

print("\nPer-component breakdown:")
print(f"{'Component':22s} {'Failures':>8s} {'Downtime(h)':>12s} {'Share':>7s} {'MTTR(h)':>8s}")
for comp, d in sorted(per_comp.items(), key=lambda kv: -kv[1]["downtime"]):
    d["share"] = d["downtime"] / total_downtime
    d["mttr"] = d["downtime"] / d["failures"]
    print(f"{comp:22s} {d['failures']:8d} {d['downtime']:12d} "
          f"{d['share']*100:6.1f}% {d['mttr']:8.2f}")

# inter-failure gaps (time between consecutive failure events)
gaps = [EVENTS[0][2]] + [EVENTS[i][2] - EVENTS[i-1][3] for i in range(1, n_failures)]
print(f"\nObserved up-intervals between failures (h): {gaps}")
print(f"Mean of observed up-intervals: {sum(gaps)/len(gaps):.2f} h (last interval "
      f"{OPERATING_HOURS-EVENTS[-1][3]} h is right-censored, not included)")

# Student-specific additional view: service-level downtime assuming redundancy
# masks single-node failures (only single points of failure take the service down)
spof_downtime = sum(ev[4] for ev in EVENTS if ev[1] in SINGLE_POINTS_OF_FAILURE)
spof_availability = (OPERATING_HOURS - spof_downtime) / OPERATING_HOURS
print(f"\n[Additional view] If redundancy masks A/B and DB failures, only SPOF events "
      f"(LB + Network) cause user-visible outage:")
print(f"  SPOF downtime = {spof_downtime} h  ->  service availability = "
      f"{spof_availability:.4f} = {spof_availability*100:.2f} %")

# --------------------------------------------------------------------------
# Part B - Reliability Block Diagram
# --------------------------------------------------------------------------
section("PART B - RELIABILITY BLOCK DIAGRAM")
R = RELIABILITY
R_app = 1 - (1 - R["Application Server A"]) * (1 - R["Application Server B"])
R_db = 1 - (1 - R["Primary Database"]) * (1 - R["Standby Database"])
R_lb = R["Load Balancer"]
R_net = R["Campus Network"]
R_sys = R_lb * R_app * R_db * R_net

print(f"Application (A OR B, parallel): 1 - (1-{R['Application Server A']})(1-{R['Application Server B']}) "
      f"= 1 - {(1-R['Application Server A'])*(1-R['Application Server B']):.6f} = {R_app:.6f}")
print(f"Database (Primary OR Standby) : 1 - (1-{R['Primary Database']})(1-{R['Standby Database']}) "
      f"= 1 - {(1-R['Primary Database'])*(1-R['Standby Database']):.6f} = {R_db:.6f}")
print(f"Load Balancer (single)        : {R_lb}")
print(f"Campus Network (single)       : {R_net}")
print(f"Overall (series)              : {R_lb} x {R_app:.6f} x {R_db:.6f} x {R_net} = {R_sys:.6f}")
print(f"System unreliability          : {1-R_sys:.6f}")
print("Contribution of each subsystem to unreliability (1 - R_i):")
for name, r in [("Load Balancer", R_lb), ("Campus Network", R_net),
                ("Application subsystem", R_app), ("Database subsystem", R_db)]:
    print(f"  {name:24s} {1-r:.6f}")

# What-if: no redundancy at all
R_no_red = R_lb * R["Application Server A"] * R["Primary Database"] * R_net
print(f"What-if (no redundancy, A + Primary only): {R_no_red:.6f}")
# What-if: redundant load balancer + dual network uplink
R_lb2 = 1 - (1 - R_lb) ** 2
R_net2 = 1 - (1 - R_net) ** 2
R_sys_improved = R_lb2 * R_app * R_db * R_net2
print(f"What-if (redundant LB and dual network uplink): {R_sys_improved:.6f}")

# --------------------------------------------------------------------------
# Part C - FMEA
# --------------------------------------------------------------------------
section("PART C - FMEA (RPN = S x O x D)")
# (component, failure mode, effect, cause, S, O, D, mitigation)
FMEA = [
    ("Load Balancer", "Node crash / hardware or process failure",
     "Complete portal outage; no requests reach A/B (single point of failure)",
     "Hardware fault, OS/firmware bug, config error during change",
     10, 3, 3, "Deploy active-passive LB pair with VRRP/keepalived; health-check alert"),
    ("Load Balancer", "Misrouting / faulty health check marks healthy servers down",
     "Traffic sent to failed node or all traffic to one server; latency and errors",
     "Wrong health-check path, stale configuration",
     6, 3, 6, "Config validation in CI, synthetic end-to-end probes, canary config rollout"),
    ("Application Server A", "Crash / OOM / process hang",
     "Capacity halved; B takes full load; risk of overload at peak registration",
     "Memory leak, unhandled exception, JVM/GC pressure",
     5, 6, 3, "Auto-restart (systemd/k8s liveness), memory limits, add 3rd node (N+2)"),
    ("Application Server A", "Degraded performance (slow responses, not down)",
     "Timeouts for students; LB may not detect because node still answers",
     "Thread-pool exhaustion, slow DB queries, disk full",
     6, 5, 7, "Latency-based health checks, APM tracing, p95 latency alerting"),
    ("Application Server B", "Crash / OOM / process hang",
     "Capacity halved; A takes full load",
     "Memory leak, unhandled exception",
     5, 5, 3, "Same as A: liveness probes, autoscaling, N+2 capacity"),
    ("Application Server B", "Failed deployment / bad release",
     "Errors on B only; inconsistent behaviour between nodes",
     "Untested release, missing rollback",
     6, 4, 5, "Blue-green / canary deploys, automated rollback, smoke tests"),
    ("Primary Database", "Node failure / crash",
     "Writes blocked until failover; registrations lost or delayed (3 events, 13 h)",
     "Disk failure, OOM, long-running lock, storage saturation",
     9, 7, 4, "Automatic failover (Patroni/Always On) < 60 s; tested failover drills"),
    ("Primary Database", "Data corruption / replication lag",
     "Wrong or stale data shown to students; failover to inconsistent standby",
     "Replication break, unsafe schema migration, lost WAL",
     9, 3, 8, "Replication-lag monitoring, checksums, PITR backups, restore tests"),
    ("Standby Database", "Standby not synchronised / fails to promote",
     "Redundancy silently lost; primary failure becomes full outage",
     "Replication stopped, insufficient resources, untested promotion",
     8, 3, 9, "Alert on replication lag > threshold; monthly failover rehearsal"),
    ("Standby Database", "Standby node hardware failure",
     "No effect while primary healthy; hidden loss of redundancy",
     "Disk / power failure",
     4, 2, 6, "Hardware monitoring, standby health in dashboard"),
    ("Campus Network", "Core switch / uplink outage",
     "Complete outage for on-campus users; remote users may also lose access (SPOF)",
     "Switch failure, fibre cut, power outage, misconfiguration",
     10, 4, 2, "Dual uplinks with separate paths, redundant core switches, UPS"),
    ("Campus Network", "Congestion / DDoS during registration peak",
     "Slow or dropped requests exactly when demand is highest",
     "Traffic spike, DDoS, lack of QoS",
     7, 5, 5, "Rate limiting, DDoS protection, QoS for portal traffic, CDN for static assets"),
]

rows = []
for comp, mode, effect, cause, s, o, d, mit in FMEA:
    rpn = s * o * d
    rows.append((comp, mode, effect, cause, s, o, d, rpn, mit))

print(f"{'#':>2s} {'Component':22s} {'Failure mode':48s} {'S':>2s} {'O':>2s} {'D':>2s} {'RPN':>4s}")
for i, r in enumerate(rows, 1):
    print(f"{i:2d} {r[0]:22s} {r[1][:48]:48s} {r[4]:2d} {r[5]:2d} {r[6]:2d} {r[7]:4d}")

top3 = sorted(rows, key=lambda r: -r[7])[:3]
print("\nThree highest-RPN failure modes:")
for r in top3:
    print(f"  RPN {r[7]:3d}  {r[0]} - {r[1]}")

# --------------------------------------------------------------------------
# Part D - Fault Tree Analysis
# --------------------------------------------------------------------------
section("PART D - FAULT TREE ANALYSIS")
Q = {k: round(1 - v, 6) for k, v in RELIABILITY.items()}   # basic-event failure probabilities
Q_app = Q["Application Server A"] * Q["Application Server B"]       # AND gate
Q_db = Q["Primary Database"] * Q["Standby Database"]                # AND gate
# OR gate (exact): 1 - product of (1 - q_i)
Q_top = 1 - (1 - Q["Load Balancer"]) * (1 - Q["Campus Network"]) * (1 - Q_app) * (1 - Q_db)
Q_top_rare = Q["Load Balancer"] + Q["Campus Network"] + Q_app + Q_db  # rare-event approximation

print("Basic event probabilities q = 1 - R:")
for k, v in Q.items():
    print(f"  {k:22s} {v:.4f}")
print(f"\nG1 Application subsystem unavailable (AND): {Q['Application Server A']} x "
      f"{Q['Application Server B']} = {Q_app:.6f}")
print(f"G2 Database subsystem unavailable (AND)   : {Q['Primary Database']} x "
      f"{Q['Standby Database']} = {Q_db:.6f}")
print(f"TOP Complete service unavailable (OR)     : 1 - (1-{Q['Load Balancer']})(1-{Q['Campus Network']})"
      f"(1-{Q_app:.6f})(1-{Q_db:.6f}) = {Q_top:.6f}")
print(f"Rare-event approximation (sum)            : {Q_top_rare:.6f}")
print(f"Consistency with RBD: 1 - R_sys = {1-R_sys:.6f}  ->  {'OK' if abs(Q_top-(1-R_sys))<1e-12 else 'MISMATCH'}")
print("Minimal cut sets: {LB}, {Network}, {AppA, AppB}, {PrimaryDB, StandbyDB}")
print("Importance (share of TOP probability, rare-event):")
for name, q in [("Load Balancer", Q["Load Balancer"]), ("Campus Network", Q["Campus Network"]),
                ("App A AND App B", Q_app), ("Primary AND Standby DB", Q_db)]:
    print(f"  {name:24s} {q:.6f}  ({q/Q_top_rare*100:5.1f} %)")

# --------------------------------------------------------------------------
# Export: Excel workbook
# --------------------------------------------------------------------------
section("EXPORT")
wb = Workbook()
hdr_fill = PatternFill("solid", fgColor="1F3864")
hdr_font = Font(bold=True, color="FFFFFF")
thin = Side(style="thin", color="999999")
border = Border(left=thin, right=thin, top=thin, bottom=thin)


def write_table(ws, start_row, headers, data, widths=None):
    for j, h in enumerate(headers, 1):
        c = ws.cell(row=start_row, column=j, value=h)
        c.fill, c.font, c.border = hdr_fill, hdr_font, border
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for i, row in enumerate(data, start_row + 1):
        for j, v in enumerate(row, 1):
            c = ws.cell(row=i, column=j, value=v)
            c.border = border
            c.alignment = Alignment(wrap_text=True, vertical="top")
    if widths:
        for j, w in enumerate(widths, 1):
            ws.column_dimensions[ws.cell(row=1, column=j).column_letter].width = w
    return start_row + len(data) + 2


# Sheet 1: raw data + Part A with live formulas
ws = wb.active
ws.title = "PartA_Metrics"
ws["A1"] = "Assignment 1 - Reliability metrics (Zhumabayev Magzhan, 255408)"
ws["A1"].font = Font(bold=True, size=13)
r = write_table(ws, 3, ["Event", "Component", "Failure time (h)", "Repair completed (h)", "Duration (h)"],
                [list(ev) for ev in EVENTS], [8, 24, 18, 20, 14])
first, last = 4, 3 + n_failures
ws.cell(row=r, column=1, value="Operating period (h)"); ws.cell(row=r, column=2, value=OPERATING_HOURS)
ws.cell(row=r+1, column=1, value="Number of failures"); ws.cell(row=r+1, column=2, value=f"=COUNT(A{first}:A{last})")
ws.cell(row=r+2, column=1, value="Total downtime (h)"); ws.cell(row=r+2, column=2, value=f"=SUM(E{first}:E{last})")
ws.cell(row=r+3, column=1, value="Total uptime (h)"); ws.cell(row=r+3, column=2, value=f"=B{r}-B{r+2}")
ws.cell(row=r+4, column=1, value="Availability"); ws.cell(row=r+4, column=2, value=f"=B{r+3}/B{r}")
ws.cell(row=r+4, column=2).number_format = "0.00%"
ws.cell(row=r+5, column=1, value="MTTF (h)"); ws.cell(row=r+5, column=2, value=f"=B{r+3}/B{r+1}")
ws.cell(row=r+6, column=1, value="MTTR (h)"); ws.cell(row=r+6, column=2, value=f"=B{r+2}/B{r+1}")
ws.cell(row=r+7, column=1, value="MTBF (h)"); ws.cell(row=r+7, column=2, value=f"=B{r}/B{r+1}")
ws.cell(row=r+8, column=1, value="Check: MTTF+MTTR"); ws.cell(row=r+8, column=2, value=f"=B{r+5}+B{r+6}")
for i in range(r, r + 9):
    ws.cell(row=i, column=1).font = Font(bold=True)
    ws.cell(row=i, column=1).border = border; ws.cell(row=i, column=2).border = border
r += 10
write_table(ws, r, ["Component", "Failures", "Downtime (h)", "Share of downtime", "Component MTTR (h)"],
            [[c, d["failures"], d["downtime"], round(d["share"], 4), round(d["mttr"], 2)]
             for c, d in sorted(per_comp.items(), key=lambda kv: -kv[1]["downtime"])])

# Sheet 2: RBD
ws = wb.create_sheet("PartB_RBD")
write_table(ws, 1, ["Component", "Reliability"], [[k, v] for k, v in RELIABILITY.items()], [26, 14])
rb = 10
write_table(ws, rb, ["Subsystem", "Configuration", "Formula", "Reliability"], [
    ["Application", "A OR B (parallel)", "1-(1-RA)(1-RB)", f"=1-(1-B3)*(1-B4)"],
    ["Database", "Primary OR Standby (parallel)", "1-(1-RP)(1-RS)", f"=1-(1-B5)*(1-B6)"],
    ["Load Balancer", "Single (series)", "R_LB", "=B2"],
    ["Campus Network", "Single (series)", "R_NET", "=B7"],
    ["Overall", "Series of subsystems", "R_LB x R_APP x R_DB x R_NET", f"=D{rb+3}*D{rb+1}*D{rb+2}*D{rb+4}"],
], [16, 30, 30, 14])
for i in range(rb + 1, rb + 6):
    ws.cell(row=i, column=4).number_format = "0.000000"

# Sheet 3: FMEA
ws = wb.create_sheet("PartC_FMEA")
data = []
for i, rr in enumerate(rows, 2):
    data.append([rr[0], rr[1], rr[2], rr[3], rr[4], rr[5], rr[6], f"=E{i}*F{i}*G{i}", rr[8]])
write_table(ws, 1, ["Component", "Failure mode", "Effect on AITU service", "Cause", "S", "O", "D", "RPN", "Mitigation"],
            data, [20, 34, 40, 34, 4, 4, 4, 6, 44])

# Sheet 4: FTA
ws = wb.create_sheet("PartD_FTA")
write_table(ws, 1, ["Basic event", "q = 1 - R"], [[k, round(v, 4)] for k, v in Q.items()], [30, 14])
write_table(ws, 10, ["Intermediate / top event", "Gate", "Basic events", "Formula", "Probability"], [
    ["Application subsystem unavailable", "AND", "App A fails, App B fails", "qA*qB", "=B3*B4"],
    ["Database subsystem unavailable", "AND", "Primary fails, Standby fails", "qP*qS", "=B5*B6"],
    ["Complete service unavailable (TOP)", "OR", "LB fails, Network fails, G1, G2",
     "1-(1-qLB)(1-qNET)(1-G1)(1-G2)", "=1-(1-B2)*(1-B7)*(1-E11)*(1-E12)"],
], [36, 8, 34, 30, 14])
for i in range(11, 14):
    ws.cell(row=i, column=5).number_format = "0.000000"

wb.save("reliability_calculations.xlsx")
print("Saved reliability_calculations.xlsx")

# --------------------------------------------------------------------------
# Diagrams
# --------------------------------------------------------------------------
NAVY, LIGHT, RED, GREEN = "#1F3864", "#DCE6F2", "#C00000", "#2E7D32"


def block(ax, x, y, w, h, text, color=LIGHT, edge=NAVY, fs=9):
    ax.add_patch(FancyBboxPatch((x, y), w, h, boxstyle="round,pad=0.02,rounding_size=0.08",
                                fc=color, ec=edge, lw=1.6))
    ax.text(x + w / 2, y + h / 2, text, ha="center", va="center", fontsize=fs, wrap=True)


# ---- RBD
fig, ax = plt.subplots(figsize=(12, 4.6))
ax.set_xlim(0, 12); ax.set_ylim(0, 4.6); ax.axis("off")
bw, bh = 1.7, 0.7
# LB
block(ax, 1.0, 1.95, bw, bh, f"Load Balancer\nR = {R_lb:.3f}")
# App parallel
block(ax, 4.0, 2.9, bw, bh, f"App Server A\nR = {R['Application Server A']:.3f}")
block(ax, 4.0, 1.0, bw, bh, f"App Server B\nR = {R['Application Server B']:.3f}")
# DB parallel
block(ax, 7.0, 2.9, bw, bh, f"Primary DB\nR = {R['Primary Database']:.3f}")
block(ax, 7.0, 1.0, bw, bh, f"Standby DB\nR = {R['Standby Database']:.3f}")
# Network
block(ax, 10.0, 1.95, bw, bh, f"Campus Network\nR = {R_net:.3f}")
lw = 1.6
def line(x1, y1, x2, y2): ax.plot([x1, x2], [y1, y2], color=NAVY, lw=lw)
# in/out
line(0.3, 2.3, 1.0, 2.3); line(11.7, 2.3, 12.0, 2.3)
ax.text(0.15, 2.3, "IN", ha="right", va="center", fontsize=9)
# LB -> split
line(2.7, 2.3, 3.4, 2.3); line(3.4, 1.35, 3.4, 3.25)
line(3.4, 3.25, 4.0, 3.25); line(3.4, 1.35, 4.0, 1.35)
# app -> join -> split
line(5.7, 3.25, 6.4, 3.25); line(5.7, 1.35, 6.4, 1.35); line(6.4, 1.35, 6.4, 3.25)
line(6.4, 3.25, 7.0, 3.25); line(6.4, 1.35, 7.0, 1.35)
# db -> join -> network
line(8.7, 3.25, 9.4, 3.25); line(8.7, 1.35, 9.4, 1.35); line(9.4, 1.35, 9.4, 3.25)
line(9.4, 2.3, 10.0, 2.3)
# labels
ax.text(1.85, 1.6, "series (SPOF)", ha="center", fontsize=8, color=RED)
ax.text(10.85, 1.6, "series (SPOF)", ha="center", fontsize=8, color=RED)
ax.text(4.85, 4.15, f"Application subsystem (parallel)\nR = 1-(1-0.970)^2 = {R_app:.4f}", ha="center", fontsize=8.5, color=GREEN)
ax.text(7.85, 4.15, f"Database subsystem (parallel)\nR = 1-(0.020)(0.010) = {R_db:.4f}", ha="center", fontsize=8.5, color=GREEN)
ax.text(6, 0.35, f"R_system = {R_lb} x {R_app:.4f} x {R_db:.4f} x {R_net} = {R_sys:.4f}   "
        f"(unreliability {1-R_sys:.4f})", ha="center", fontsize=10, fontweight="bold", color=NAVY)
ax.set_title("Reliability Block Diagram - AITU student portal (synthetic)", fontsize=12, color=NAVY, pad=4)
fig.savefig("rbd_diagram.png", dpi=200, bbox_inches="tight")
plt.close(fig)
print("Saved rbd_diagram.png")


# ---- FTA
def or_gate(ax, x, y, s=0.32):
    # classic OR shape: curved bottom, pointed top
    import numpy as np
    t = np.linspace(-1, 1, 40)
    top_x = x + s * t
    top_y = y + s * 1.1 * (1 - t ** 2) ** 0.5 * 1.0 + s * 0.3 * (1 - np.abs(t))
    bot_x = x + s * t[::-1]
    bot_y = y + s * 0.35 * (1 - t[::-1] ** 2) ** 0.5
    pts = list(zip(top_x, top_y)) + list(zip(bot_x, bot_y))
    ax.add_patch(Polygon(pts, closed=True, fc="#FCE4D6", ec=NAVY, lw=1.5))
    ax.text(x, y + s * 0.55, "OR", ha="center", va="center", fontsize=8, fontweight="bold")


def and_gate(ax, x, y, s=0.32):
    import numpy as np
    t = np.linspace(-1, 1, 40)
    pts = [(x - s, y), (x + s, y)] + [(x + s * tt, y + s * 0.5 + s * 1.0 * (1 - tt ** 2) ** 0.5) for tt in t[::-1]]
    ax.add_patch(Polygon(pts, closed=True, fc="#E2EFDA", ec=NAVY, lw=1.5))
    ax.text(x, y + s * 0.55, "AND", ha="center", va="center", fontsize=8, fontweight="bold")


def basic(ax, x, y, text, r=0.42):
    ax.add_patch(Circle((x, y), r, fc="#FFF2CC", ec=NAVY, lw=1.5))
    ax.text(x, y, text, ha="center", va="center", fontsize=7.5)


fig, ax = plt.subplots(figsize=(12, 7.2))
ax.set_xlim(0, 12); ax.set_ylim(0, 7.2); ax.axis("off")
# top event
block(ax, 3.7, 6.1, 4.6, 0.85,
      f"TOP: AITU student portal unavailable during\nregistration / examination period\nq_TOP = {Q_top:.4f}",
      color="#F8CBAD", fs=9)
or_gate(ax, 6.0, 5.2)
line(6.0, 5.2 + 0.32 * 1.4, 6.0, 6.1)
# horizontal bus
line(1.5, 4.7, 10.5, 4.7); line(6.0, 4.7, 6.0, 5.2)
xs = [1.5, 4.5, 7.5, 10.5]
for x in xs: line(x, 4.1, x, 4.7)
# LB basic event & network basic event
basic(ax, 1.5, 3.55, f"E1\nLoad Balancer\nfailure\nq={Q['Load Balancer']:.3f}", r=0.62)
basic(ax, 10.5, 3.55, f"E2\nCampus\nNetwork failure\nq={Q['Campus Network']:.3f}", r=0.62)
# intermediate events
block(ax, 3.3, 3.45, 2.4, 0.7, f"G1: Application subsystem\nunavailable  q={Q_app:.4f}", fs=8.2)
block(ax, 6.3, 3.45, 2.4, 0.7, f"G2: Database subsystem\nunavailable  q={Q_db:.4f}", fs=8.2)
and_gate(ax, 4.5, 2.55); and_gate(ax, 7.5, 2.55)
line(4.5, 2.55 + 0.48, 4.5, 3.45); line(7.5, 2.55 + 0.48, 7.5, 3.45)
# app basic events
line(3.6, 2.1, 5.4, 2.1); line(4.5, 2.1, 4.5, 2.55); line(3.6, 1.6, 3.6, 2.1); line(5.4, 1.6, 5.4, 2.1)
basic(ax, 3.6, 1.05, f"E3\nApp Server A\nfails\nq={Q['Application Server A']:.3f}", r=0.55)
basic(ax, 5.4, 1.05, f"E4\nApp Server B\nfails\nq={Q['Application Server B']:.3f}", r=0.55)
# db basic events
line(6.6, 2.1, 8.4, 2.1); line(7.5, 2.1, 7.5, 2.55); line(6.6, 1.6, 6.6, 2.1); line(8.4, 1.6, 8.4, 2.1)
basic(ax, 6.6, 1.05, f"E5\nPrimary DB\nfails\nq={Q['Primary Database']:.3f}", r=0.55)
basic(ax, 8.4, 1.05, f"E6\nStandby DB\nfails\nq={Q['Standby Database']:.3f}", r=0.55)
ax.text(6, 0.2, "Minimal cut sets: {E1}, {E2}, {E3,E4}, {E5,E6}    "
        f"q_TOP = 1-(1-q1)(1-q2)(1-q3q4)(1-q5q6) = {Q_top:.4f}",
        ha="center", fontsize=9.5, fontweight="bold", color=NAVY)
ax.set_title("Fault Tree Analysis - AITU student portal unavailability (synthetic)", fontsize=12, color=NAVY, pad=4)
fig.savefig("fta_diagram.png", dpi=200, bbox_inches="tight")
plt.close(fig)
print("Saved fta_diagram.png")

# ---- Downtime by component chart
fig, ax = plt.subplots(figsize=(7, 3.4))
comps = [c for c, _ in sorted(per_comp.items(), key=lambda kv: -kv[1]["downtime"])]
vals = [per_comp[c]["downtime"] for c in comps]
cols = [RED if c in ("Primary Database",) else (NAVY if c in SINGLE_POINTS_OF_FAILURE else "#7F9CC6") for c in comps]
bars = ax.barh(comps[::-1], vals[::-1], color=cols[::-1])
for b, v, c in zip(bars, vals[::-1], comps[::-1]):
    ax.text(v + 0.15, b.get_y() + b.get_height() / 2, f"{v} h ({per_comp[c]['failures']} failures)", va="center", fontsize=8.5)
ax.set_xlabel("Downtime, hours (of 27 h total)"); ax.set_xlim(0, 17)
ax.set_title("Downtime contribution by component, 720 h period", fontsize=10.5, color=NAVY)
for s in ["top", "right"]: ax.spines[s].set_visible(False)
fig.savefig("downtime_by_component.png", dpi=200, bbox_inches="tight")
plt.close(fig)
print("Saved downtime_by_component.png")

# ---- Timeline
fig, ax = plt.subplots(figsize=(11, 2.6))
order = ["Load Balancer", "Application Server A", "Application Server B", "Primary Database", "Campus Network"]
for ev in EVENTS:
    yi = order.index(ev[1])
    ax.broken_barh([(ev[2], ev[4])], (yi - 0.35, 0.7), color=RED)
    ax.text(ev[2] + ev[4] + 3, yi, f"#{ev[0]} ({ev[4]}h)", va="center", fontsize=7.5)
ax.set_yticks(range(len(order))); ax.set_yticklabels(order, fontsize=8.5)
ax.set_xlim(0, 720); ax.set_xlabel("Operating time, hours (0-720)")
ax.grid(axis="x", alpha=0.3)
ax.set_title("Failure / repair timeline (red = component in repair)", fontsize=10.5, color=NAVY)
for s in ["top", "right"]: ax.spines[s].set_visible(False)
fig.savefig("failure_timeline.png", dpi=200, bbox_inches="tight")
plt.close(fig)
print("Saved failure_timeline.png")

# results for the report generator
import json
json.dump({
    "total_downtime": total_downtime, "total_uptime": total_uptime, "availability": availability,
    "mttf": mttf, "mttr": mttr, "mtbf": mtbf, "n_failures": n_failures,
    "per_comp": per_comp, "gaps": gaps, "spof_downtime": spof_downtime, "spof_availability": spof_availability,
    "R_app": R_app, "R_db": R_db, "R_lb": R_lb, "R_net": R_net, "R_sys": R_sys,
    "R_no_red": R_no_red, "R_sys_improved": R_sys_improved,
    "fmea": rows, "top3": top3,
    "Q": Q, "Q_app": Q_app, "Q_db": Q_db, "Q_top": Q_top, "Q_top_rare": Q_top_rare,
}, open("results.json", "w"), indent=2)
print("Saved results.json")
